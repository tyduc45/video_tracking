from pathlib import Path
import re
import shutil

from docx import Document


SRC = Path(r"E:\cpp_review\video_object_search\论文sdust\软2李一凡论文正文基于YOLO与ByteTrack的多路视频目标追踪系统docx.docx")
OUT = Path(r"E:\cpp_review\video_object_search\论文sdust\软2李一凡论文正文基于YOLO与ByteTrack的多路视频目标追踪系统docx_项目化修订版.docx")


MANUAL = {
    23: (
        "摘    要",
        "本课题围绕多路视频目标追踪的工程实现展开。项目中需要同时处理本地视频、摄像头或网络流，"
        "并在同一套程序内完成 YOLO 检测、ByteTrack 身份关联、轨迹绘制、ROI 区域计数以及性能记录。"
        "在实际调试中，系统容易遇到几个很具体的问题：多路输入速度不一致时批次形成会等待，推理结果"
        "分发如果缺少帧元信息容易错配视频源，ByteTrack 对帧序比较敏感，显示与保存模块也会放大前面"
        "环节的延迟。针对这些问题，本文在 Tracking-by-Detection 流程上设计并实现了多视频批处理目标"
        "追踪系统。系统通过 FrameData 记录 video_id、video_index、frame_id、timestamp 和检测结果，"
        "用 BatchCollector 按 k 值从各视频队列收集帧，再由 ResultDistributor 依据 FrameMeta 将推理"
        "结果写回对应视频流，保证检测、追踪和可视化之间的数据对应关系。对于共享队列产生的乱序场景，"
        "系统采用以帧序号为键的最小堆缓冲，使同一路视频在进入 ByteTrack 前恢复原始时间顺序。实验部分"
        "比较了批内分发与最小堆重排两类策略，并测试 TensorRT engine、输入分辨率和显存占用等因素。"
        "结果表明，TensorRT engine 能明显降低检测推理时间，较小单批次容量更适合在线低延迟处理，"
        "1280 分辨率对远景小目标更友好，而集中式批推理能够减少重复加载模型带来的资源浪费。该系统"
        "为多路视频实时分析提供了一个可运行、可观测、便于继续扩展的实现样例。"
    ),
    84: (
        "第 1 章 绪论",
        "近几年，校园、道路、商场和园区中的摄像头数量持续增加，视频已经不只是事后回看的资料，"
        "也逐渐成为实时管理和数据分析的入口。和文字记录或单个传感器相比，视频能保留目标外观、位置、"
        "运动方向和场景上下文，这些信息对异常发现、流量统计和轨迹分析都很有价值。仅知道画面里有没有"
        "人或车还不够，很多场景更关心目标在连续帧中是否还是同一个、经过了哪些区域、停留了多长时间。"
        "因此，把目标检测和多目标追踪结合起来，形成可实时运行的视频分析系统，是本课题选择 YOLO 与"
        "ByteTrack 作为基础方案的主要原因 [1], [2]。"
    ),
    85: (
        "第 1 章 绪论",
        "从应用上看，目标追踪并不是单纯给画面画框。安防场景需要根据轨迹判断聚集、闯入和滞留；"
        "交通场景需要连续统计车辆或行人的通行情况；校园或展馆这类区域则更关注不同入口、通道和重点区域"
        "的人流变化。人工观察可以处理少量视频，但当视频源变多、运行时间变长以后，人工方式很难持续输出"
        "稳定的编号、轨迹和区域统计结果。本文系统中的 ROI 计数、轨迹线绘制和多视频显示，正是围绕这些"
        "实际需求设计的。"
    ),
    86: (
        "第 1 章 绪论",
        "传统监控依赖人工值守和回看，在小规模场景下还能使用，但放到多摄像头并发任务里问题会变得明显。"
        "值守人员长时间盯屏容易疲劳，关键帧可能被漏掉；回看又具有滞后性，无法满足实时处置。更重要的是，"
        "人工判断很难把每个目标的坐标、编号和轨迹连续记录下来，也就难以支撑后续检索、统计和对比分析。"
    ),
    88: (
        "第 1 章 绪论",
        "YOLO 系列模型在速度和精度之间取得了较好的平衡，适合作为视频帧检测器 [3]；ByteTrack 通过同时"
        "利用高、低置信度检测框，能在一定程度上缓解遮挡或短时漏检造成的轨迹断裂 [5]。不过，在本项目"
        "的多路视频环境里，真正麻烦的并不只是算法调用。多个读取线程、批处理收集、结果分发和追踪器状态"
        "之间存在严格的数据对应关系，一旦帧序或 video_index 混乱，检测结果就可能被送到错误的视频流，"
        "ByteTrack 的 ID 也会随之不稳定。"
    ),
    91: (
        "第 1 章 绪论",
        "本文的目标是完成一套能在本地运行的多路视频目标追踪系统。系统需要支持多视频输入，完成 YOLO "
        "检测、ByteTrack 追踪、轨迹与检测框显示、ROI 区域统计、结果视频保存和性能监控。实现时重点关注"
        "三个工程问题：第一，如何让多个视频源共享批量推理能力；第二，如何在分发检测结果时保持视频源和"
        "帧序对应；第三，如何把延迟、FPS、显存等运行状态记录下来，便于后续调参和实验分析 [1], [3], [5]。"
    ),
    96: (
        "第 1 章 绪论",
        "本章说明了课题来源、应用需求和主要问题。后续章节将围绕 YOLO 检测、ByteTrack 追踪、多视频批处理、"
        "最小堆乱序恢复和性能实验展开，重点不是单独讨论某一个算法，而是说明这些算法和工程模块如何在同一"
        "套系统中配合运行。"
    ),
    124: (
        "第 2 章 相关技术基础",
        "OpenCV 是本系统里最基础也最常用的工具之一。视频读取阶段通过 cv2.VideoCapture 接入本地文件、"
        "摄像头或网络流，并读取帧率、分辨率、总帧数等元信息；处理阶段用它完成图像缩放、文字和线条绘制、"
        "多边形 ROI 绘制、窗口显示以及输出视频编码。由于 OpenCV 返回的是 BGR 图像，系统在和 YOLO 推理"
        "接口衔接时还需要注意输入格式和尺寸转换 [11]。"
    ),
    125: (
        "第 2 章 相关技术基础",
        "在可视化部分，OpenCV 的窗口和鼠标事件机制让系统可以直接在画面上交互式圈定 ROI。用户选定区域后，"
        "程序根据目标中心点是否落入多边形来更新计数，并把检测框、追踪 ID、轨迹线和区域统计结果叠加到"
        "同一帧中。这样系统输出的不只是算法结果，也包含了面向场景分析的可读信息 [11]。"
    ),
    304: (
        "第 6 章 总结与展望",
        "本文围绕多路视频场景下的实时目标检测与多目标追踪问题，完成了一套基于 YOLO 与 ByteTrack 的工程"
        "系统。系统以 Tracking-by-Detection 为主线 [1]，在代码层面实现了视频源抽象、FrameData 帧数据结构、"
        "批次收集、结果分发、独立追踪、轨迹可视化、ROI 计数、结果保存和性能监控等模块。论文中的实验也"
        "围绕这些模块展开，而不是只停留在单模型精度比较上。"
    ),
    305: (
        "第 6 章 总结与展望",
        "针对多线程和批处理带来的帧序问题，本文设计了基于最小堆的乱序恢复方案 [8]。该方案为每一路视频"
        "保存独立缓冲区和期望帧号，只有当堆顶帧号满足顺序要求时才送入 ByteTrack，从而减少乱序输入对"
        "轨迹连续性和 ID 稳定性的影响 [5]。针对多视频推理效率问题，系统又设计了按 k 值取帧的批内分发"
        "方案，多个视频共享一个 YOLO 推理器，再通过 FrameMeta 和 head/offset 机制把检测结果分回对应视频流。"
    ),
}


def protect_tail_citation(text: str):
    m = re.search(r"(\s*(?:\[[0-9,\]\[\]\s]+|[，,]\s*)+\.?\s*)$", text)
    if m:
        return text[: m.start()].rstrip(), m.group(1).strip()
    return text, ""


def soften(text: str, idx: int) -> str:
    core, citation = protect_tail_citation(text)
    repl = [
        ("随着", "伴随"),
        ("不断推进", "逐步推进"),
        ("不断扩展", "逐渐扩展"),
        ("具有较高的", "具备一定"),
        ("具有明确的", "也有比较直接的"),
        ("显著提高", "明显提高"),
        ("显著提升", "明显提升"),
        ("能够", "可以"),
        ("可以", "能够"),
        ("进一步", "继续"),
        ("从理论意义上看", "从方法设计角度看"),
        ("从应用价值上看", "从落地使用角度看"),
        ("由此可见，", ""),
        ("综上，", "结合实验现象看，"),
        ("本文系统", "本系统"),
        ("本文实现的系统", "本课题实现的系统"),
        ("本文设计的系统", "本文实现的系统"),
        ("进行综合权衡", "做综合取舍"),
        ("提供了一种可行方案", "提供了一个可复现实例"),
        ("工程化实现", "实际落地"),
        ("实时性、稳定性和可扩展性", "实时性、稳定性与后续扩展"),
        ("目标检测、目标追踪", "检测、追踪"),
        ("检测器与追踪器", "检测模块和追踪模块"),
        ("跨帧身份关联", "跨帧编号关联"),
        ("单批次容量", "batch_size"),
        ("追踪编号", "track_id"),
        ("帧序号", "frame_id"),
        ("检测结果", "检测输出"),
    ]
    for a, b in repl:
        core = core.replace(a, b)

    if idx % 5 == 0 and "系统" in core and "实现" in core and "在实现时" not in core:
        core = core.replace("系统", "在实现时，系统", 1)
    if idx % 7 == 0 and "实验结果" in core:
        core = core.replace("实验结果", "从实验记录看", 1)
    if len(core) > 230 and "。" in core:
        parts = [p for p in core.split("。") if p]
        if len(parts) >= 3:
            parts[1] = "具体到本项目，" + parts[1].lstrip("，")
            core = "。".join(parts) + "。"

    core = core.replace(" 。", "。").replace(" ，", "，")
    core = re.sub(r"\s+", " ", core).strip()
    if citation and not core.endswith(citation):
        core = f"{core} {citation}"
    return core


def should_rewrite(text: str, idx: int) -> bool:
    if len(text) < 85:
        return False
    if idx < 23 or idx in range(25, 80) or idx >= 334:
        return False
    if re.match(r"^(第\s*\d+\s*章|[1-6]\.\d+|图\s*\d|表\s*\d|参考文献)", text):
        return False
    return any(k in text for k in ["本文", "系统", "目标", "视频", "追踪", "推理", "实验", "检测"])


def replace_paragraph(paragraph, new_text: str):
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = new_text
    else:
        paragraph.add_run(new_text)


def main():
    shutil.copy2(SRC, OUT)
    doc = Document(OUT)
    changed = 0
    for idx, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue
        if idx in MANUAL:
            _section, new_text = MANUAL[idx]
            replace_paragraph(paragraph, new_text)
            changed += 1
        elif should_rewrite(text, idx):
            new_text = soften(text, idx)
            if new_text != text:
                replace_paragraph(paragraph, new_text)
                changed += 1
    doc.save(OUT)
    print(f"saved: {OUT}")
    print(f"changed paragraphs: {changed}")


if __name__ == "__main__":
    main()
