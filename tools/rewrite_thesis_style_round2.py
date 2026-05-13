from pathlib import Path
import re
import shutil

from docx import Document


BASE = Path(r"E:\cpp_review\video_object_search\论文sdust\软2李一凡论文正文基于YOLO与ByteTrack的多路视频目标追踪系统docx_项目化修订版 .docx")
USER = Path(r"E:\cpp_review\video_object_search\论文sdust\软2李一凡论文正文基于YOLO与ByteTrack的多路视频目标追踪系统docx_项目化修订版  - 副本.docx")
OUT = Path(r"E:\cpp_review\video_object_search\论文sdust\软2李一凡论文正文基于YOLO与ByteTrack的多路视频目标追踪系统docx_按个人语气续改版.docx")


def paragraph_texts(doc):
    return [p.text.strip() for p in doc.paragraphs]


def replace_paragraph(paragraph, new_text: str):
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = new_text
    else:
        paragraph.add_run(new_text)


def find_after(texts, heading):
    for i, text in enumerate(texts):
        if text == heading and i + 1 < len(texts):
            return texts[i + 1]
    return None


def replace_after_heading(doc, heading, new_text):
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == heading and i + 1 < len(doc.paragraphs):
            replace_paragraph(doc.paragraphs[i + 1], new_text)
            return True
    return False


def citation_tail(text):
    m = re.search(r"(\s*(?:\[[0-9,\]\[\]\s]+|[，,]\s*)+\.?\s*)$", text)
    if m:
        return text[: m.start()].rstrip(), m.group(1).strip()
    return text, ""


def style_rewrite(text, idx):
    body, cite = citation_tail(text)

    pairs = [
        ("本文系统", "本系统"),
        ("本文实现的系统", "本系统"),
        ("本文设计的系统", "本系统"),
        ("本文的研究目的在于", "本课题主要想完成的是"),
        ("本文的目标是", "本课题的目标是"),
        ("本文围绕", "作者围绕"),
        ("本文采用", "本系统采用"),
        ("本文将", "作者将"),
        ("本文提出", "作者设计"),
        ("本文进一步", "作者进一步"),
        ("本研究希望解决", "作者在开发过程中主要解决"),
        ("具体而言，", "具体到本项目，"),
        ("从理论意义上看", "从方法设计角度看"),
        ("从应用价值上看", "从实际使用角度看"),
        ("因此，", "因此在本项目中，"),
        ("由此可见，", ""),
        ("需要注意的是，", "这里需要说明的是，"),
        ("在实际部署中，", "在作者实际测试和部署时，"),
        ("在实际应用中，", "在真实使用场景里，"),
        ("随着", "伴随"),
        ("不断", "持续"),
        ("进一步", "继续"),
        ("显著", "明显"),
        ("具有较强的工程适用性", "在工程上比较好用"),
        ("具有较好的工程应用价值", "有比较明确的工程价值"),
        ("提供了一种可行方案", "提供了一个可以继续扩展的实现基础"),
        ("追踪编号", "track_id"),
        ("帧序号", "frame_id"),
        ("单批次容量", "batch_size"),
        ("检测结果", "检测输出"),
        ("检测器与追踪器", "检测模块和追踪模块"),
        ("多目标追踪系统", "MOT 系统"),
        ("实时视频多目标追踪", "实时 MOT"),
        ("跨帧身份关联", "跨帧编号关联"),
    ]
    for a, b in pairs:
        body = body.replace(a, b)

    if len(body) > 170 and idx % 4 == 0 and "本系统" in body:
        body = "在本系统中，" + body.lstrip("在本系统中，")
    if len(body) > 190 and idx % 6 == 0 and "。" in body:
        parts = [p for p in body.split("。") if p]
        if len(parts) >= 3 and not parts[1].startswith("作者"):
            parts[1] = "作者在实现时重点关注这一点：" + parts[1]
            body = "。".join(parts) + "。"
    if len(body) > 220 and idx % 9 == 0 and "：" not in body[:80]:
        body = body.replace("，", "，一方面，", 1)
        body = body.replace("；", "；另一方面，", 1)

    body = body.replace("能够能够", "能够").replace("可以可以", "可以")
    body = body.replace("问题问题", "问题").replace("以及。", "。")
    body = re.sub(r"\s+", " ", body).strip()
    if cite:
        body = f"{body} {cite}"
    return body


def should_touch(text, idx):
    if len(text) < 95:
        return False
    if idx < 23 or idx >= 335:
        return False
    if "\t" in text:
        return False
    if re.match(r"^(第\s*\d+\s*章|[1-6]\.\d+|图\s*\d|表\s*\d|Algorithm|参考文献)", text):
        return False
    return any(k in text for k in ("本文", "系统", "视频", "追踪", "检测", "推理", "实验", "ByteTrack", "YOLO"))


ROUND2_MANUAL = {
    "摘    要": (
        "本课题围绕多路视频场景下的实时目标检测与多目标追踪问题展开。项目周期中，作者独立完成了一套"
        "基于 YOLO 与 ByteTrack 的工程系统，系统能够接入多路本地视频、摄像头或网络流，并完成检测、"
        "track_id 关联、轨迹绘制、ROI 区域计数、结果保存和性能监控等功能。开发过程中遇到的主要问题"
        "并不是单纯调用模型，而是多线程读取、批处理推理、结果分发和追踪器状态之间的协同：多路输入速度"
        "不一致时会影响 batch 形成；检测输出如果没有绑定 FrameMeta，就可能分发到错误视频源；ByteTrack "
        "对 frame_id 顺序比较敏感，帧序错乱会直接影响 ID 稳定性。针对这些问题，作者设计了两类处理策略："
        "一类是基于最小堆的乱序帧恢复机制，另一类是按 k 值取帧的批内分发机制。系统通过 FrameData 保存"
        "video_id、video_index、frame_id、timestamp 和检测输出，通过 BatchCollector 收集批次，再由 "
        "ResultDistributor 使用 head/offset 机制把结果写回各视频队列。实验部分围绕乱序恢复策略、batch_size、"
        "TensorRT engine、输入分辨率和 GPU 显存占用进行测试。实验结果说明，TensorRT engine 能明显降低"
        "推理延迟，较小 batch_size 更适合在线实时处理，1280 分辨率更利于远景小目标识别，集中式批推理"
        "相比多线程并行流水线能够减少一定 GPU 显存占用。整体来看，本系统不是一个单独算法 demo，而是一个"
        "包含读取、推理、追踪、显示、计数、监控和实验脚本的多视频目标追踪工程实现。"
    ),
    "1.1 研究背景": (
        "在本项目周期中，作者首先关注的是多路视频监控在真实使用中的压力。校园、道路、商场和园区的视频源"
        "数量持续增加，人工回看虽然简单，但很难同时回答“目标是谁、从哪里来、经过哪里、停留多久”这些连续"
        "性问题。视频帧本身包含目标外观、位置、运动方向和场景上下文，如果只做单帧检测，很多时间信息会被"
        "浪费。因此，本课题选择把 YOLO 检测和 ByteTrack 追踪组合起来，在实时画面中同时输出检测框、track_id、"
        "轨迹线和区域统计结果，为后续的视频目标搜索和场景分析打基础 [1], [2]。"
    ),
    "1.2 研究目的与意义": (
        "本课题的目的不是单独比较某一个检测模型的精度，而是做出一套可以运行、可以观察、可以做实验的多路视频"
        "追踪系统。具体目标包括：1. 支持多个视频源同时输入；2. 使用 YOLO 完成逐帧检测；3. 使用 ByteTrack "
        "维护同一目标的 track_id；4. 支持轨迹绘制、ROI 区域计数和输出视频保存；5. 记录端到端延迟、FPS 和"
        "显存占用等实验数据。通过这些模块，作者希望验证多线程读取、批处理推理和乱序恢复策略对系统实时性与"
        "稳定性的实际影响 [1], [3], [5]。"
    ),
    "2.1 多目标追踪任务概述": (
        "多目标追踪（Multiple Object Tracking，MOT）可以理解为“在连续视频帧中给同一目标保持同一个编号”的"
        "任务 [1], [2]。检测模型只回答当前帧中有哪些目标以及目标在哪里，而追踪算法还要回答下一帧中的哪个框"
        "和上一帧属于同一个人或同一辆车。在本系统中，YOLO 负责输出目标框、类别和置信度，ByteTrack 负责在"
        "这些检测输出之间建立时间关系，并给目标分配 track_id。后续 ROI 计数、轨迹绘制和目标检索都依赖这个"
        "编号是否稳定。"
    ),
    "3.1 系统需求分析": (
        "本系统面向多路视频目标检测与追踪场景，作者在需求分析阶段把任务拆成了几个必须落地的功能：1. 视频源"
        "能够来自本地文件、摄像头或网络流；2. 每一帧必须带有 video_id、video_index、frame_id 和 timestamp；"
        "3. YOLO 推理输出要能准确回到原视频队列；4. 每一路视频需要维护独立 ByteTrack 状态，避免多视频之间"
        "互相污染 track_id；5. 显示和保存模块需要把检测框、轨迹、ROI 计数和性能信息一起输出。相比单视频脚本，"
        "多路视频系统更容易在队列等待、批次收集和结果分发处出问题，因此需求分析中也把实时性、稳定性和资源"
        "占用作为重点指标 [1], [2]。"
    ),
}


SECTION_REPLACEMENTS = {
    ("5.7 实验结果综合分析", "5.8 本章小结"): [
        "综合实验结果，作者得到以下几点结论。第一，ByteTrack 对单路视频内部的 frame_id 顺序比较敏感，乱序输入会破坏轨迹连续性，因此在多线程或弱网场景下必须显式处理帧序。第二，batch_size 并不是越大越好。较大 batch 能提高 GPU 吞吐，但在线追踪更关心端到端延迟，在本实验条件下 batch_size 为 4 更适合实时显示。第三，TensorRT engine 相比 pt 模型带来了明显推理加速，FPS 从 36.68 提升到 250.96，单帧推理延迟从 27.26 ms 降到 3.98 ms。第四，1280 分辨率对远景小目标更友好，但运行耗时也会增加。第五，集中式批推理比多线程并行流水线节省约 22 MiB GPU 显存，虽然差距不算夸张，但说明共享推理实例确实能减少重复资源占用。",
        "从工程部署角度看，本系统更适合把 TensorRT engine、batch_size 为 4 和批内分发策略作为默认实时配置；当任务更关注低尾部延迟，并且能够接受视频源配额不完全均衡时，可以参考最小堆重排策略的共享队列调度思路；当任务更关注远景小目标识别时，可以切换到 1280 输入分辨率，但需要接受额外耗时。上述实验从功能正确性、时序稳定性、推理速度、分辨率影响和资源占用多个角度验证了系统设计的有效性。"
    ],
}


def cleanup_text(text):
    replacements = {
        "纳入MOT": "纳入 MOT",
        "输入纳入MOT": "输入纳入 MOT",
        "保存video_id": "保存 video_id",
        "通过FrameData": "通过 FrameData",
        "track_id和": "track_id 和",
        "batch_size为": "batch_size 为",
        "batch_size更": "batch_size 更",
        "batch_size或": "batch_size 或",
        "大batch_size": "大 batch_size",
        "动态batch_size": "动态 batch_size",
        "较小batch_size": "较小 batch_size",
        "具体到本项目，作者在开发过程中主要解决三个方面的问题。具体到本项目，第一": "作者在开发过程中主要解决三个方面的问题。第一",
        "在本项目周期中，作者首先关注的是多路视频监控在真实使用中的压力。作者在实现时重点关注这一点：": "在本项目周期中，作者首先关注的是多路视频监控在真实使用中的压力。",
        "系统能够接入": "系统可以接入",
    }
    for a, b in replacements.items():
        text = text.replace(a, b)
    return text


USER_SECTION6 = [
    ("6.1 工作总结", [
        "在本项目周期中，作者围绕多路视频场景下的实时目标检测与多目标追踪问题，独立研发了一套基于 YOLO 与 ByteTrack 的工程系统。以下是核心工作总结：",
        "1. 系统以 Tracking-by-Detection 为主线 [1]，在代码层面实现了视频源抽象、FrameData 帧数据结构、批次收集、结果分发、独立追踪、轨迹可视化、ROI 计数、结果保存和性能监控等模块。",
        "2. 作者设计了两种乱序恢复方案，用于解决多线程静态条件和弱网环境下可能出现的 frame_id 错乱问题 [8]。",
        "3. 本系统实现了多视频实时显示和交互式 ROI 区域计数功能。",
        "4. 系统实现了采用探针式设计的性能监控模块并提供可视化支持，可以记录单帧端到端延迟并绘制实时延迟曲线，为系统调试和实验分析提供数据支持。",
        "5. 作者利用 Windows 下的命令行脚本语言实现了自动化实验运行脚本，提高了测试效率。",
        "6. 作者设计了针对不同 batch_size 以及不同输入分辨率的模型 .engine 文件导出机制，使系统兼具灵活性和高性能推理能力。",
        "7. 作者自主设计实验，并客观分析了系统在不同分辨率、不同乱序恢复策略下的性能差异，同时通过实验数据证明本系统比多线程并行流水线更加节省 GPU 显存。"
    ]),
    ("6.2 主要创新与亮点", [
        "本文工作的主要亮点体现在以下几个方面。",
        "第一，将帧序一致性问题纳入多目标追踪系统设计，提出并实现了基于最小堆的乱序帧恢复机制。该机制针对 ByteTrack 对时间顺序敏感的特点，在追踪器之前增加轻量级时序修正环节，使系统能够在异步读取和共享批处理条件下维持单视频内部顺序 [5], [8]。",
        "第二，设计了面向多路视频的批处理推理与结果分发机制。系统通过配额分配机制为不同视频分配单批次内帧数，通过批次收集器按序收集多路视频帧，并通过结果分发器根据单帧元数据将检测输出写回对应视频队列。该方案既发挥了 YOLO 批量推理的吞吐优势，又避免了多视频结果混淆和追踪状态污染 [9], [10]。",
        "第三，实现了较完整的工程化可视化与监控功能。显示系统可以根据视频源数量自动调整布局，用户能够在实时画面中观察检测框、追踪 ID 和轨迹线，也可以通过鼠标绘制多边形 ROI，并基于多边形包裹判定算法实时统计区域内目标数量。同时，通过性能监控器和性能显示窗口对运行延迟进行实时观测，使系统从单纯算法调用扩展为可运行、可观察、可实验的多视频目标追踪平台。",
        "第四，作者分析并比较了两种方法对系统的影响。实验结果显示，基于最小堆的乱序帧恢复方法在多数场景下优于按 k 值取帧的批内分发方案；同时作者发现，后者由于源码设计中存在串行瓶颈，使得按 k 值取帧方案仍存在继续并行优化的研究空间。"
    ]),
]


def replace_section_paragraphs(doc, heading, new_paras, next_heading):
    texts = paragraph_texts(doc)
    start = next(i for i, t in enumerate(texts) if t == heading)
    end = next(i for i in range(start + 1, len(texts)) if texts[i] == next_heading)
    old_slots = list(range(start + 1, end))
    for pos, text in zip(old_slots, new_paras):
        replace_paragraph(doc.paragraphs[pos], text)
    for pos in old_slots[len(new_paras):]:
        replace_paragraph(doc.paragraphs[pos], "")
    if len(new_paras) > len(old_slots):
        next_heading_para = doc.paragraphs[end]
        for text in new_paras[len(old_slots):]:
            next_heading_para.insert_paragraph_before(text)


def main():
    shutil.copy2(BASE, OUT)
    doc = Document(OUT)
    user_doc = Document(USER)
    user_texts = paragraph_texts(user_doc)

    changed = 0

    # Preserve the user's already rewritten chapter summaries where the large
    # image-bearing document has the corresponding headings.
    summary_map = {
        "1.3 本章小结": find_after(user_texts, "1.3 本章小结"),
        "2.9 本章小结": find_after(user_texts, "2.9 本章小结"),
        "3.8 本章小结": find_after(user_texts, "3.4 本章小结") or find_after(user_texts, "3.8 本章小结"),
        "6.5 本章小结": find_after(user_texts, "6.5 本章小结"),
    }
    for heading, text in summary_map.items():
        if text and replace_after_heading(doc, heading, text):
            changed += 1

    # Manual high-impact rewrites by section heading.
    for heading, text in ROUND2_MANUAL.items():
        if replace_after_heading(doc, heading, text):
            changed += 1

    for (heading, next_heading), paras in SECTION_REPLACEMENTS.items():
        replace_section_paragraphs(doc, heading, paras, next_heading)
        changed += len(paras)

    replace_section_paragraphs(doc, "6.1 工作总结", USER_SECTION6[0][1], "6.2 主要创新与亮点")
    replace_section_paragraphs(doc, "6.2 主要创新与亮点", USER_SECTION6[1][1], "6.3 不足之处")
    changed += sum(len(x[1]) for x in USER_SECTION6)

    # Lightly restyle the remaining long paragraphs.
    for idx, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if should_touch(text, idx):
            new_text = style_rewrite(text, idx)
            if new_text != text:
                replace_paragraph(p, new_text)
                changed += 1

    for p in doc.paragraphs:
        cleaned = cleanup_text(p.text)
        if cleaned != p.text:
            replace_paragraph(p, cleaned)
            changed += 1

    doc.save(OUT)
    print(f"saved: {OUT}")
    print(f"changed-ish paragraphs: {changed}")


if __name__ == "__main__":
    main()
